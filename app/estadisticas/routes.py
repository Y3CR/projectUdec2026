from flask import render_template, redirect, url_for, flash, request, jsonify, send_file
from flask_login import login_required, current_user
from functools import wraps
from datetime import datetime, timedelta
from io import BytesIO
import pytz
from app import db
from app.estadisticas import bp
from app.models import (Solicitud, Recurso, Espacio, User,
                        RegistroAcceso, DispositivoRFID, PermisoAcceso)

BOGOTA_TZ = pytz.timezone('America/Bogota')


def admin_o_operador(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated:
            return redirect(url_for('auth.login'))
        if current_user.get_role_name() not in ['administrador', 'operador']:
            flash('Acceso denegado.', 'danger')
            return redirect(url_for('admin.dashboard'))
        return f(*args, **kwargs)
    return decorated


def _rango_fechas(periodo, fecha_desde_str='', fecha_hasta_str=''):
    """Retorna (fecha_inicio, fecha_fin) según el período."""
    ahora = datetime.now(BOGOTA_TZ).replace(tzinfo=None)

    if periodo == 'personalizado':
        fecha_ini = None
        fecha_fin = None
        try:
            if fecha_desde_str:
                fecha_ini = datetime.strptime(fecha_desde_str, '%Y-%m-%d')
        except ValueError:
            pass
        try:
            if fecha_hasta_str:
                fecha_fin = (datetime.strptime(fecha_hasta_str, '%Y-%m-%d')
                             + timedelta(days=1))
        except ValueError:
            pass
        return fecha_ini, fecha_fin

    if periodo == 'semana':
        return ahora - timedelta(days=7), ahora
    elif periodo == 'mes':
        return ahora - timedelta(days=30), ahora
    elif periodo == 'trimestre':
        return ahora - timedelta(days=90), ahora
    elif periodo == 'anio':
        return ahora - timedelta(days=365), ahora
    return None, None



# ── Dashboard principal ────────────────────────────────────────────────────────

@bp.route('/dashboard')
@login_required
@admin_o_operador
def dashboard():
    periodo = request.args.get('periodo', 'mes')
    fecha_desde_str = request.args.get('fecha_desde', '')
    fecha_hasta_str = request.args.get('fecha_hasta', '')

    fecha_ini, fecha_fin = _rango_fechas(
        periodo, fecha_desde_str, fecha_hasta_str
    )

    q_base = Solicitud.query
    if fecha_ini:
        q_base = q_base.filter(Solicitud.fecha_creacion >= fecha_ini)
    if fecha_fin:
        q_base = q_base.filter(Solicitud.fecha_creacion <= fecha_fin)

    kpis = {
        'total_solicitudes': q_base.count(),
        'aprobadas': q_base.filter(
            Solicitud.estado.in_(['aprobada', 'entregado', 'devuelta'])
        ).count(),
        'rechazadas': q_base.filter(
            Solicitud.estado == 'rechazada'
        ).count(),
        'en_uso': Solicitud.query.filter_by(
            estado='entregado', tipo='recurso'
        ).count(),
        'recursos_disponibles': Recurso.query.filter_by(
            estado='disponible'
        ).count(),
        'total_recursos': Recurso.query.count(),
        'total_usuarios_activos': User.query.filter_by(activo=True).count(),
        'accesos_rfid': (
            RegistroAcceso.query.filter(
                RegistroAcceso.fecha_evento >= fecha_ini
            ).count() if fecha_ini
            else RegistroAcceso.query.count()
        ),
    }

    from app.models import Role
    roles_ids = [r.id for r in Role.query.filter(
        Role.name.in_(['docente', 'estudiante'])
    ).all()]
    usuarios = User.query.filter(
        User.activo == True,
        User.role_id.in_(roles_ids)
    ).order_by(User.nombre).all()

    espacios = Espacio.query.order_by(Espacio.nombre).all()
    recursos_lista = Recurso.query.order_by(Recurso.nombre).all()

    return render_template('estadisticas/dashboard.html',
                           kpis=kpis,
                           periodo=periodo,
                           fecha_desde=fecha_desde_str,
                           fecha_hasta=fecha_hasta_str,
                           usuarios=usuarios,
                           espacios=espacios,
                           recursos_lista=recursos_lista)
# ── APIs de datos para Chart.js ────────────────────────────────────────────────

@bp.route('/api/solicitudes-por-estado')
@login_required
@admin_o_operador
def api_solicitudes_estado():
    periodo = request.args.get('periodo', 'mes')
    fecha_ini, fecha_fin = _rango_fechas(periodo)

    q = Solicitud.query
    if fecha_ini:
        q = q.filter(Solicitud.fecha_creacion >= fecha_ini)
    if fecha_fin:
        q = q.filter(Solicitud.fecha_creacion <= fecha_fin)

    estados = ['pendiente', 'aprobada', 'entregado', 'devuelta', 'rechazada']
    datos = []
    for e in estados:
        datos.append(q.filter(Solicitud.estado == e).count())

    return jsonify({
        'labels': ['Pendiente', 'Aprobada', 'Entregado', 'Devuelta', 'Rechazada'],
        'data': datos,
        'colors': ['#6c757d', '#0d6efd', '#198754', '#0dcaf0', '#dc3545']
    })



@bp.route('/api/solicitudes-por-mes')
@login_required
@admin_o_operador
def api_solicitudes_mes():
    periodo = request.args.get('periodo', 'mes')
    ahora = datetime.now(BOGOTA_TZ).replace(tzinfo=None)
    labels, aprobadas, rechazadas, totales = [], [], [], []

    if periodo == 'semana':
        # Últimos 7 días agrupados por día
        for i in range(6, -1, -1):
            fecha = ahora - timedelta(days=i)
            ini = fecha.replace(hour=0, minute=0, second=0, microsecond=0)
            fin = fecha.replace(hour=23, minute=59, second=59)
            labels.append(fecha.strftime('%a %d/%m'))
            q = Solicitud.query.filter(
                Solicitud.fecha_creacion >= ini,
                Solicitud.fecha_creacion <= fin
            )
            totales.append(q.count())
            aprobadas.append(q.filter(
                Solicitud.estado.in_(['aprobada', 'entregado', 'devuelta'])
            ).count())
            rechazadas.append(q.filter(
                Solicitud.estado == 'rechazada'
            ).count())

    elif periodo == 'mes':
        # Últimos 30 días agrupados por día
        for i in range(29, -1, -1):
            fecha = ahora - timedelta(days=i)
            ini = fecha.replace(hour=0, minute=0, second=0, microsecond=0)
            fin = fecha.replace(hour=23, minute=59, second=59)
            labels.append(fecha.strftime('%d/%m'))
            q = Solicitud.query.filter(
                Solicitud.fecha_creacion >= ini,
                Solicitud.fecha_creacion <= fin
            )
            totales.append(q.count())
            aprobadas.append(q.filter(
                Solicitud.estado.in_(['aprobada', 'entregado', 'devuelta'])
            ).count())
            rechazadas.append(q.filter(
                Solicitud.estado == 'rechazada'
            ).count())

    elif periodo == 'trimestre':
        # Últimas 13 semanas agrupadas por semana
        for i in range(12, -1, -1):
            fin = ahora - timedelta(weeks=i)
            ini = fin - timedelta(days=6)
            labels.append(f"S {fin.strftime('%d/%m')}")
            q = Solicitud.query.filter(
                Solicitud.fecha_creacion >= ini,
                Solicitud.fecha_creacion <= fin
            )
            totales.append(q.count())
            aprobadas.append(q.filter(
                Solicitud.estado.in_(['aprobada', 'entregado', 'devuelta'])
            ).count())
            rechazadas.append(q.filter(
                Solicitud.estado == 'rechazada'
            ).count())

    else:
        # anio o personalizado: últimos 12 meses agrupados por mes
        for i in range(11, -1, -1):
            fecha = ahora - timedelta(days=30 * i)
            mes_ini = fecha.replace(day=1, hour=0, minute=0, second=0)
            if i == 0:
                mes_fin = ahora
            else:
                siguiente = fecha.replace(day=28) + timedelta(days=4)
                mes_fin = siguiente.replace(day=1)
            labels.append(fecha.strftime('%b %Y'))
            q = Solicitud.query.filter(
                Solicitud.fecha_creacion >= mes_ini,
                Solicitud.fecha_creacion < mes_fin
            )
            totales.append(q.count())
            aprobadas.append(q.filter(
                Solicitud.estado.in_(['aprobada', 'entregado', 'devuelta'])
            ).count())
            rechazadas.append(q.filter(
                Solicitud.estado == 'rechazada'
            ).count())

    return jsonify({
        'labels': labels,
        'datasets': [
            {'label': 'Total', 'data': totales, 'color': '#1b4332'},
            {'label': 'Aprobadas', 'data': aprobadas, 'color': '#198754'},
            {'label': 'Rechazadas', 'data': rechazadas, 'color': '#dc3545'},
        ]
    })






@bp.route('/api/recursos-mas-solicitados')
@login_required
@admin_o_operador
def api_recursos_solicitados():
    periodo = request.args.get('periodo', 'mes')
    fecha_ini, fecha_fin = _rango_fechas(periodo)

    q = db.session.query(
        Recurso.nombre,
        db.func.count(Solicitud.id).label('total')
    ).join(Solicitud, Solicitud.recurso_id == Recurso.id)

    if fecha_ini:
        q = q.filter(Solicitud.fecha_creacion >= fecha_ini)
    if fecha_fin:
        q = q.filter(Solicitud.fecha_creacion <= fecha_fin)

    resultados = q.group_by(Recurso.id).order_by(
        db.func.count(Solicitud.id).desc()
    ).limit(10).all()

    return jsonify({
        'labels': [r.nombre for r in resultados],
        'data': [r.total for r in resultados],
        'color': '#1b4332'
    })


@bp.route('/api/espacios-mas-accedidos')
@login_required
@admin_o_operador
def api_espacios_accedidos():
    periodo = request.args.get('periodo', 'mes')
    fecha_ini, fecha_fin = _rango_fechas(periodo)

    q = db.session.query(
        Espacio.nombre,
        db.func.count(RegistroAcceso.id).label('total')
    ).join(RegistroAcceso, RegistroAcceso.espacio_id == Espacio.id)

    if fecha_ini:
        q = q.filter(RegistroAcceso.fecha_evento >= fecha_ini)
    if fecha_fin:
        q = q.filter(RegistroAcceso.fecha_evento <= fecha_fin)

    resultados = q.group_by(Espacio.id).order_by(
        db.func.count(RegistroAcceso.id).desc()
    ).limit(8).all()

    colores = ['#1b4332', '#2d6a4f', '#40916c', '#52b788',
               '#74c69d', '#95d5b2', '#b7e4c7', '#d8f3dc']

    return jsonify({
        'labels': [r.nombre for r in resultados],
        'data': [r.total for r in resultados],
        'colors': colores[:len(resultados)]
    })


@bp.route('/api/accesos-por-hora')
@login_required
@admin_o_operador
def api_accesos_hora():
    periodo = request.args.get('periodo', 'mes')
    fecha_ini, fecha_fin = _rango_fechas(periodo)

    horas = list(range(0, 24))
    permitidos = []
    denegados = []

    for h in horas:
        q_base = RegistroAcceso.query.filter(
            db.func.strftime('%H', RegistroAcceso.fecha_evento) == f'{h:02d}'
        )
        if fecha_ini:
            q_base = q_base.filter(RegistroAcceso.fecha_evento >= fecha_ini)
        if fecha_fin:
            q_base = q_base.filter(RegistroAcceso.fecha_evento <= fecha_fin)

        permitidos.append(q_base.filter_by(autorizado=True).count())
        denegados.append(q_base.filter_by(autorizado=False).count())

    return jsonify({
        'labels': [f'{h:02d}:00' for h in horas],
        'datasets': [
            {'label': 'Permitidos', 'data': permitidos, 'color': '#198754'},
            {'label': 'Denegados', 'data': denegados, 'color': '#dc3545'},
        ]
    })


@bp.route('/api/tiempo-uso-promedio')
@login_required
@admin_o_operador
def api_tiempo_uso():
    periodo = request.args.get('periodo', 'mes')
    fecha_ini, fecha_fin = _rango_fechas(periodo)

    q = db.session.query(
        Recurso.nombre,
        db.func.avg(Solicitud.tiempo_uso_minutos).label('promedio')
    ).join(Solicitud, Solicitud.recurso_id == Recurso.id).filter(
        Solicitud.tiempo_uso_minutos.isnot(None)
    )

    if fecha_ini:
        q = q.filter(Solicitud.fecha_creacion >= fecha_ini)
    if fecha_fin:
        q = q.filter(Solicitud.fecha_creacion <= fecha_fin)

    resultados = q.group_by(Recurso.id).order_by(
        db.func.avg(Solicitud.tiempo_uso_minutos).desc()
    ).limit(8).all()

    def fmt(minutos):
        if minutos is None:
            return 0
        return round(minutos / 60, 1)

    return jsonify({
        'labels': [r.nombre for r in resultados],
        'data': [fmt(r.promedio) for r in resultados],
        'color': '#b8901f'
    })


@bp.route('/api/usuarios-mas-activos')
@login_required
@admin_o_operador
def api_usuarios_activos():
    periodo = request.args.get('periodo', 'mes')
    fecha_ini, fecha_fin = _rango_fechas(periodo)

    q = db.session.query(
        User.nombre,
        User.apellido,
        db.func.count(Solicitud.id).label('total')
    ).join(Solicitud, Solicitud.usuario_id == User.id)

    if fecha_ini:
        q = q.filter(Solicitud.fecha_creacion >= fecha_ini)
    if fecha_fin:
        q = q.filter(Solicitud.fecha_creacion <= fecha_fin)

    resultados = q.group_by(User.id).order_by(
        db.func.count(Solicitud.id).desc()
    ).limit(8).all()

    return jsonify({
        'labels': [f'{r.nombre} {r.apellido}' for r in resultados],
        'data': [r.total for r in resultados],
        'color': '#0d6efd'
    })


# ── Estadística individual ─────────────────────────────────────────────────────


@bp.route('/individual')
@login_required
@admin_o_operador
def individual():
    tipo = request.args.get('tipo', 'usuario')
    item_id = request.args.get('id', 0, type=int)
    periodo = request.args.get('periodo', 'mes')
    fecha_desde_str = request.args.get('fecha_desde', '')
    fecha_hasta_str = request.args.get('fecha_hasta', '')

    fecha_ini, fecha_fin = _rango_fechas(
        periodo, fecha_desde_str, fecha_hasta_str
    )

    from app.models import Role
    roles_ids = [r.id for r in Role.query.filter(
        Role.name.in_(['docente', 'estudiante'])
    ).all()]
    usuarios = User.query.filter(
        User.activo == True,
        User.role_id.in_(roles_ids)
    ).order_by(User.nombre).all()

    espacios = Espacio.query.order_by(Espacio.nombre).all()
    recursos_lista = Recurso.query.order_by(Recurso.nombre).all()

    contexto = {
        'tipo': tipo,
        'item_id': item_id,
        'periodo': periodo,
        'fecha_desde': fecha_desde_str,
        'fecha_hasta': fecha_hasta_str,
        'usuarios': usuarios,
        'espacios': espacios,
        'recursos_lista': recursos_lista,
        'item': None,
        'datos': {}
    }

    if tipo == 'usuario' and item_id:
        usuario = User.query.get(item_id)
        if usuario:
            q = Solicitud.query.filter_by(usuario_id=item_id)
            if fecha_ini:
                q = q.filter(Solicitud.fecha_creacion >= fecha_ini)
            if fecha_fin:
                q = q.filter(Solicitud.fecha_creacion <= fecha_fin)
            todas = q.order_by(Solicitud.fecha_creacion.desc()).all()
            contexto['item'] = usuario
            contexto['datos'] = {
                'total': len(todas),
                'aprobadas': sum(1 for s in todas
                                 if s.estado in ['aprobada', 'entregado',
                                                  'devuelta']),
                'rechazadas': sum(1 for s in todas
                                  if s.estado == 'rechazada'),
                'en_uso': sum(1 for s in todas
                              if s.estado == 'entregado'),
                'tiempo_total': sum(s.tiempo_uso_minutos or 0
                                    for s in todas),
                'solicitudes': todas[:10]
            }

    elif tipo == 'recurso' and item_id:
        recurso = Recurso.query.get(item_id)
        if recurso:
            q = Solicitud.query.filter_by(recurso_id=item_id)
            if fecha_ini:
                q = q.filter(Solicitud.fecha_creacion >= fecha_ini)
            if fecha_fin:
                q = q.filter(Solicitud.fecha_creacion <= fecha_fin)
            todas = q.order_by(Solicitud.fecha_creacion.desc()).all()
            con_tiempo = [s for s in todas if s.tiempo_uso_minutos]
            tiempo_prom = (
                sum(s.tiempo_uso_minutos for s in con_tiempo)
                // len(con_tiempo) if con_tiempo else 0
            )
            contexto['item'] = recurso
            contexto['datos'] = {
                'total_solicitudes': len(todas),
                'veces_prestado': sum(1 for s in todas
                                      if s.estado in ['entregado', 'devuelta']),
                'tiempo_promedio': tiempo_prom,
                'solicitudes': todas[:10]
            }

    elif tipo == 'espacio' and item_id:
        espacio = Espacio.query.get(item_id)
        if espacio:
            q = RegistroAcceso.query.filter_by(espacio_id=item_id)
            if fecha_ini:
                q = q.filter(RegistroAcceso.fecha_evento >= fecha_ini)
            if fecha_fin:
                q = q.filter(RegistroAcceso.fecha_evento <= fecha_fin)
            todos = q.order_by(RegistroAcceso.fecha_evento.desc()).all()
            contexto['item'] = espacio
            contexto['datos'] = {
                'total_accesos': len(todos),
                'permitidos': sum(1 for a in todos if a.autorizado),
                'denegados': sum(1 for a in todos if not a.autorizado),
                'accesos': todos[:10]
            }

    return render_template('estadisticas/individual.html', **contexto)



@bp.route('/api/individual/solicitudes-por-mes')
@login_required
@admin_o_operador
def api_individual_solicitudes():
    tipo = request.args.get('tipo', 'usuario')
    item_id = request.args.get('id', 0, type=int)
    periodo = request.args.get('periodo', 'mes')
    fecha_desde_str = request.args.get('fecha_desde', '')
    fecha_hasta_str = request.args.get('fecha_hasta', '')

    ahora = datetime.now(BOGOTA_TZ).replace(tzinfo=None)
    labels = []
    datos = []

    if periodo == 'personalizado':
        # Para personalizado: agrupar por día dentro del rango
        fecha_ini, fecha_fin = _rango_fechas(
            'personalizado', fecha_desde_str, fecha_hasta_str
        )
        if not fecha_ini:
            fecha_ini = ahora - timedelta(days=30)
        if not fecha_fin:
            fecha_fin = ahora

        diff_dias = (fecha_fin - fecha_ini).days

        if diff_dias <= 31:
            # Agrupar por día
            for i in range(diff_dias + 1):
                fecha = fecha_ini + timedelta(days=i)
                dia_ini = fecha.replace(
                    hour=0, minute=0, second=0, microsecond=0
                )
                dia_fin = fecha.replace(hour=23, minute=59, second=59)
                labels.append(fecha.strftime('%d/%m'))
                datos.append(_contar_actividad(
                    tipo, item_id, dia_ini, dia_fin
                ))
        elif diff_dias <= 90:
            # Agrupar por semana
            semanas = diff_dias // 7 + 1
            for i in range(semanas):
                sem_ini = fecha_ini + timedelta(weeks=i)
                sem_fin = min(sem_ini + timedelta(days=6), fecha_fin)
                labels.append(f"S {sem_ini.strftime('%d/%m')}")
                datos.append(_contar_actividad(
                    tipo, item_id, sem_ini, sem_fin
                ))
        else:
            # Agrupar por mes
            cursor = fecha_ini.replace(
                day=1, hour=0, minute=0, second=0
            )
            while cursor <= fecha_fin:
                siguiente = cursor.replace(day=28) + timedelta(days=4)
                mes_fin = min(siguiente.replace(day=1), fecha_fin)
                labels.append(cursor.strftime('%b %Y'))
                datos.append(_contar_actividad(
                    tipo, item_id, cursor, mes_fin
                ))
                cursor = siguiente.replace(day=1)

    elif periodo == 'semana':
        for i in range(6, -1, -1):
            fecha = ahora - timedelta(days=i)
            dia_ini = fecha.replace(
                hour=0, minute=0, second=0, microsecond=0
            )
            dia_fin = fecha.replace(hour=23, minute=59, second=59)
            labels.append(fecha.strftime('%a %d/%m'))
            datos.append(_contar_actividad(tipo, item_id, dia_ini, dia_fin))

    elif periodo == 'mes':
        for i in range(29, -1, -1):
            fecha = ahora - timedelta(days=i)
            dia_ini = fecha.replace(
                hour=0, minute=0, second=0, microsecond=0
            )
            dia_fin = fecha.replace(hour=23, minute=59, second=59)
            labels.append(fecha.strftime('%d/%m'))
            datos.append(_contar_actividad(tipo, item_id, dia_ini, dia_fin))

    elif periodo == 'trimestre':
        for i in range(12, -1, -1):
            sem_fin = ahora - timedelta(weeks=i)
            sem_ini = sem_fin - timedelta(days=6)
            labels.append(f"S {sem_fin.strftime('%d/%m')}")
            datos.append(_contar_actividad(tipo, item_id, sem_ini, sem_fin))

    else:
        # anio o default: últimos 12 meses por mes
        for i in range(11, -1, -1):
            fecha = ahora - timedelta(days=30 * i)
            mes_ini = fecha.replace(day=1, hour=0, minute=0, second=0)
            if i == 0:
                mes_fin = ahora
            else:
                siguiente = fecha.replace(day=28) + timedelta(days=4)
                mes_fin = siguiente.replace(day=1)
            labels.append(fecha.strftime('%b %Y'))
            datos.append(_contar_actividad(tipo, item_id, mes_ini, mes_fin))

    return jsonify({'labels': labels, 'data': datos, 'color': '#1b4332'})    
    
    

def _contar_actividad(tipo, item_id, fecha_ini, fecha_fin):
    """Helper para contar eventos en un rango de fechas."""
    if tipo == 'usuario':
        return Solicitud.query.filter(
            Solicitud.usuario_id == item_id,
            Solicitud.fecha_creacion >= fecha_ini,
            Solicitud.fecha_creacion <= fecha_fin
        ).count()
    elif tipo == 'recurso':
        return Solicitud.query.filter(
            Solicitud.recurso_id == item_id,
            Solicitud.fecha_creacion >= fecha_ini,
            Solicitud.fecha_creacion <= fecha_fin
        ).count()
    elif tipo == 'espacio':
        return RegistroAcceso.query.filter(
            RegistroAcceso.espacio_id == item_id,
            RegistroAcceso.fecha_evento >= fecha_ini,
            RegistroAcceso.fecha_evento <= fecha_fin
        ).count()
    return 0

# ── Exportar Excel del dashboard ───────────────────────────────────────────────

@bp.route('/exportar/excel')
@login_required
@admin_o_operador
def exportar_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    periodo = request.args.get('periodo', 'mes')
    fecha_ini, fecha_fin = _rango_fechas(periodo)

    wb = openpyxl.Workbook()
    verde_fill = PatternFill(start_color='1b4332', end_color='1b4332', fill_type='solid')
    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    data_font = Font(name='Calibri', size=10)
    center = Alignment(horizontal='center', vertical='center')
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def escribir_hoja(ws, titulo, headers, filas):
        ws.merge_cells(f'A1:{get_column_letter(len(headers))}1')
        ws['A1'] = titulo
        ws['A1'].font = Font(name='Calibri', bold=True, size=13, color='1b4332')
        ws['A1'].alignment = center
        ws.row_dimensions[1].height = 22

        for col, h in enumerate(headers, 1):
            c = ws.cell(row=2, column=col, value=h)
            c.font = header_font
            c.fill = verde_fill
            c.alignment = center
            c.border = border
        ws.row_dimensions[2].height = 18

        for row_num, fila in enumerate(filas, 3):
            alt = PatternFill(
                start_color='f0f7f3' if row_num % 2 == 0 else 'FFFFFF',
                end_color='f0f7f3' if row_num % 2 == 0 else 'FFFFFF',
                fill_type='solid'
            )
            for col, val in enumerate(fila, 1):
                c = ws.cell(row=row_num, column=col, value=val)
                c.font = data_font
                c.alignment = center
                c.border = border
                c.fill = alt

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20

    # Hoja 1: Solicitudes
    ws1 = wb.active
    ws1.title = 'Solicitudes'
    q = Solicitud.query
    if fecha_ini:
        q = q.filter(Solicitud.fecha_creacion >= fecha_ini)
    if fecha_fin:
        q = q.filter(Solicitud.fecha_creacion <= fecha_fin)
    solicitudes = q.order_by(Solicitud.fecha_creacion.desc()).all()
    filas1 = []
    for s in solicitudes:
        horas = (s.tiempo_uso_minutos or 0) // 60
        mins = (s.tiempo_uso_minutos or 0) % 60
        filas1.append([
            s.id,
            f'{s.usuario.nombre} {s.usuario.apellido}',
            s.usuario.get_role_name().capitalize(),
            s.tipo.upper(),
            s.get_item_nombre(),
            s.fecha_creacion.strftime('%d/%m/%Y %H:%M'),
            s.fecha_inicio.strftime('%d/%m/%Y %H:%M') if s.fecha_inicio else '—',
            s.fecha_fin_efectiva().strftime('%d/%m/%Y %H:%M') if s.fecha_fin_efectiva() else '—',
            s.estado.capitalize(),
            f'{horas}h {mins}min' if s.tiempo_uso_minutos else '—'
        ])
    escribir_hoja(ws1, f'Solicitudes — Período: {periodo}',
                  ['#', 'Usuario', 'Rol', 'Tipo', 'Item', 'Creada',
                   'Inicio', 'Fin estimado', 'Estado', 'Tiempo uso'],
                  filas1)

    # Hoja 2: Recursos más solicitados
    ws2 = wb.create_sheet('Recursos top')
    res_q = db.session.query(
        Recurso.nombre, Recurso.serial,
        db.func.count(Solicitud.id).label('total')
    ).join(Solicitud, Solicitud.recurso_id == Recurso.id)
    if fecha_ini:
        res_q = res_q.filter(Solicitud.fecha_creacion >= fecha_ini)
    res_top = res_q.group_by(Recurso.id).order_by(
        db.func.count(Solicitud.id).desc()
    ).all()
    escribir_hoja(ws2, 'Recursos más solicitados',
                  ['Recurso', 'Serial', 'Total solicitudes'],
                  [(r.nombre, r.serial or '—', r.total) for r in res_top])

    # Hoja 3: Accesos RFID
    ws3 = wb.create_sheet('Accesos RFID')
    acc_q = RegistroAcceso.query
    if fecha_ini:
        acc_q = acc_q.filter(RegistroAcceso.fecha_evento >= fecha_ini)
    if fecha_fin:
        acc_q = acc_q.filter(RegistroAcceso.fecha_evento <= fecha_fin)
    accesos = acc_q.order_by(RegistroAcceso.fecha_evento.desc()).limit(500).all()
    filas3 = []
    for a in accesos:
        filas3.append([
            a.fecha_evento.strftime('%d/%m/%Y %H:%M:%S'),
            f'{a.usuario.nombre} {a.usuario.apellido}' if a.usuario else 'Desconocido',
            a.espacio.nombre if a.espacio else '—',
            a.dispositivo.codigo if a.dispositivo else '—',
            a.tipo_evento.capitalize(),
            '✅ Permitido' if a.autorizado else '❌ Denegado',
            a.motivo_denegacion or '—'
        ])
    escribir_hoja(ws3, 'Registros de acceso RFID',
                  ['Fecha/Hora', 'Usuario', 'Espacio', 'Dispositivo',
                   'Tipo', 'Resultado', 'Motivo'],
                  filas3)

    # Hoja 4: Usuarios más activos
    ws4 = wb.create_sheet('Usuarios activos')
    usr_q = db.session.query(
        User.nombre, User.apellido,
        db.func.count(Solicitud.id).label('total')
    ).join(Solicitud, Solicitud.usuario_id == User.id)
    if fecha_ini:
        usr_q = usr_q.filter(Solicitud.fecha_creacion >= fecha_ini)
    usr_top = usr_q.group_by(User.id).order_by(
        db.func.count(Solicitud.id).desc()
    ).all()
    escribir_hoja(ws4, 'Usuarios más activos',
                  ['Nombre', 'Apellido', 'Total solicitudes'],
                  [(u.nombre, u.apellido, u.total) for u in usr_top])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    nombre = f"estadisticas_{periodo}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(buf,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=nombre)
    


@bp.route('/exportar/csv-ml', methods=['GET', 'POST'])
@login_required
@admin_o_operador
def exportar_csv_ml():
    import csv
    from io import StringIO

    # Definición de todos los campos disponibles
    CAMPOS_DISPONIBLES = {
        'solicitud_id':             'ID de solicitud',
        'usuario_id':               'ID de usuario',
        'recurso_id':               'ID de recurso',
        'anio_creacion':            'Año de creación',
        'mes_creacion':             'Mes de creación (1-12)',
        'dia_semana_creacion':      'Día de la semana (0=lun)',
        'hora_creacion':            'Hora de creación (0-23)',
        'trimestre_creacion':       'Trimestre (1-4)',
        'rol_usuario':              'Rol del usuario',
        'tiene_tarjeta_rfid':       'Tiene tarjeta RFID (0/1)',
        'total_prestamos_previos':  'Préstamos previos',
        'total_devoluciones_a_tiempo': 'Devoluciones a tiempo',
        'total_demoras_previas':    'Demoras previas',
        'ratio_demoras':            'Ratio de demoras (0-1)',
        'categoria_recurso':        'Categoría del recurso',
        'veces_prestado_recurso':   'Veces prestado el recurso',
        'dias_prestamo_solicitados':'Días de préstamo solicitados',
        'tiene_motivo':             'Tiene motivo (0/1)',
        'longitud_motivo':          'Palabras en el motivo',
        'fue_aprobada':             'TARGET: Fue aprobada (0/1)',
        'fue_entregada':            'TARGET: Fue entregada (0/1)',
        'tuvo_demora':              'TARGET: Tuvo demora (0/1)',
        'tiempo_uso_real_minutos':  'TARGET: Tiempo uso (min)',
        'tiempo_uso_real_horas':    'TARGET: Tiempo uso (hrs)',
        'dias_retraso':             'TARGET: Días de retraso',
        'estado_devolucion':        'TARGET: Estado devolución',
    }

    # Si es GET, mostrar el formulario
    if request.method == 'GET':
        from app.models import Role
        return render_template(
            'estadisticas/csv_ml_config.html',
            campos=CAMPOS_DISPONIBLES
        )

    # Si es POST, generar el CSV con los parámetros seleccionados
    campos_sel = request.form.getlist('campos')
    fecha_desde_str = request.form.get('fecha_desde', '').strip()
    fecha_hasta_str = request.form.get('fecha_hasta', '').strip()

    if not campos_sel:
        campos_sel = list(CAMPOS_DISPONIBLES.keys())

    fecha_ini = None
    fecha_fin = None
    try:
        if fecha_desde_str:
            fecha_ini = datetime.strptime(fecha_desde_str, '%Y-%m-%d')
        if fecha_hasta_str:
            fecha_fin = datetime.strptime(fecha_hasta_str, '%Y-%m-%d') \
                        + timedelta(days=1)
    except ValueError:
        pass

    q = Solicitud.query.filter(Solicitud.tipo == 'recurso')
    if fecha_ini:
        q = q.filter(Solicitud.fecha_creacion >= fecha_ini)
    if fecha_fin:
        q = q.filter(Solicitud.fecha_creacion <= fecha_fin)
    solicitudes = q.order_by(Solicitud.fecha_creacion.asc()).all()

    # Caché de métricas
    cache_usuario = {}

    def _metricas(usuario_id, hasta):
        key = (usuario_id, hasta.date() if hasta else None)
        if key in cache_usuario:
            return cache_usuario[key]
        anteriores = Solicitud.query.filter(
            Solicitud.usuario_id == usuario_id,
            Solicitud.tipo == 'recurso',
            Solicitud.fecha_creacion < hasta
        ).all() if hasta else []
        total = len(anteriores)
        demoras = sum(1 for s in anteriores if (
            s.estado == 'en_demora' or (
                s.estado == 'devuelta'
                and s.fecha_devolucion_real
                and s.fecha_fin_efectiva()
                and s.fecha_devolucion_real > s.fecha_fin_efectiva()
            )
        ))
        a_tiempo = sum(1 for s in anteriores if (
            s.estado == 'devuelta'
            and s.fecha_devolucion_real
            and s.fecha_fin_efectiva()
            and s.fecha_devolucion_real <= s.fecha_fin_efectiva()
        ))
        ratio = round(demoras / total, 4) if total > 0 else 0.0
        result = (total, a_tiempo, demoras, ratio)
        cache_usuario[key] = result
        return result

    cache_recurso = {}

    def _veces(recurso_id, hasta):
        key = (recurso_id, hasta.date() if hasta else None)
        if key in cache_recurso:
            return cache_recurso[key]
        n = Solicitud.query.filter(
            Solicitud.recurso_id == recurso_id,
            Solicitud.tipo == 'recurso',
            Solicitud.fecha_creacion < hasta,
            Solicitud.estado.in_(['entregado', 'devuelta', 'en_demora'])
        ).count() if hasta else 0
        cache_recurso[key] = n
        return n

    # Construir filas
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow([CAMPOS_DISPONIBLES[c] for c in campos_sel])

    for s in solicitudes:
        fc = s.fecha_creacion
        u = s.usuario
        r = s.recurso

        total_prev, a_tiempo_prev, demoras_prev, ratio_dem = _metricas(
            s.usuario_id, fc
        )

        if s.fecha_fin and fc:
            dias_sol = max(0, (s.fecha_fin - fc).days)
        else:
            dias_sol = None

        motivo = s.motivo or ''

        if s.fecha_devolucion_real and s.fecha_fin_efectiva():
            diff = s.fecha_devolucion_real - s.fecha_fin_efectiva()
            dias_retraso = round(diff.total_seconds() / 86400, 2)
        else:
            dias_retraso = None

        tiempo_min = s.tiempo_uso_minutos

        todos_valores = {
            'solicitud_id':             s.id,
            'usuario_id':               s.usuario_id,
            'recurso_id':               s.recurso_id,
            'anio_creacion':            fc.year,
            'mes_creacion':             fc.month,
            'dia_semana_creacion':      fc.weekday(),
            'hora_creacion':            fc.hour,
            'trimestre_creacion':       (fc.month - 1) // 3 + 1,
            'rol_usuario':              u.get_role_name() if u else '',
            'tiene_tarjeta_rfid':       1 if (u and u.uid_tarjeta) else 0,
            'total_prestamos_previos':  total_prev,
            'total_devoluciones_a_tiempo': a_tiempo_prev,
            'total_demoras_previas':    demoras_prev,
            'ratio_demoras':            ratio_dem,
            'categoria_recurso':        r.categoria.nombre if r else '',
            'veces_prestado_recurso':   _veces(s.recurso_id, fc),
            'dias_prestamo_solicitados': dias_sol,
            'tiene_motivo':             1 if motivo.strip() else 0,
            'longitud_motivo':          len(motivo.strip().split()) if motivo.strip() else 0,
            'fue_aprobada':             1 if s.estado in [
                'aprobada', 'entregado', 'devuelta', 'en_demora'
            ] else 0,
            'fue_entregada':            1 if s.estado in [
                'entregado', 'devuelta', 'en_demora'
            ] else 0,
            'tuvo_demora':              1 if (
                s.estado == 'en_demora' or (
                    s.estado == 'devuelta'
                    and s.fecha_devolucion_real
                    and s.fecha_fin_efectiva()
                    and s.fecha_devolucion_real > s.fecha_fin_efectiva()
                )
            ) else 0,
            'tiempo_uso_real_minutos':  tiempo_min,
            'tiempo_uso_real_horas':    round(tiempo_min / 60, 2) if tiempo_min else None,
            'dias_retraso':             dias_retraso,
            'estado_devolucion':        s.estado_devolucion or '',
        }

        writer.writerow([todos_valores[c] for c in campos_sel])

    output.seek(0)
    buf = BytesIO(output.getvalue().encode('utf-8-sig'))
    nombre = f"ml_prestamos_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    return send_file(
        buf,
        mimetype='text/csv',
        as_attachment=True,
        download_name=nombre
    )